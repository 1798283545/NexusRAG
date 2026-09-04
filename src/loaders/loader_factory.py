"""文档加载器工厂模块。

根据文件扩展名返回对应的文档加载器，屏蔽各格式的具体解析细节。
调用方式::

    docs = LoaderFactory.get_loader(file_path).load()

支持的格式: PDF、Word(.docx)、Markdown、TXT。
"""

from __future__ import annotations

import logging
import os
from abc import ABC, abstractmethod
from datetime import datetime
from typing import Dict, List, Optional, Type

from langchain_community.document_loaders import TextLoader as _LangchainTextLoader
from langchain_core.documents import Document

logger = logging.getLogger(__name__)

#: 工厂支持的文件扩展名
SUPPORTED_EXTENSIONS: tuple[str, ...] = (".pdf", ".docx", ".md", ".txt")


class LoaderError(Exception):
    """文档加载过程中发生错误的基类异常。"""


class UnsupportedFormatError(LoaderError, ValueError):
    """文件格式不受支持时抛出的异常。"""

    def __init__(self, file_path: str, extension: str) -> None:
        super().__init__(
            f"不支持的文件格式 '{extension}'（文件: {file_path}）。"
            f"支持的格式: {', '.join(SUPPORTED_EXTENSIONS)}"
        )


class BaseLoader(ABC):
    """所有文档加载器的抽象基类。

    定义统一的 ``load`` 接口与元数据提取逻辑；各格式加载器只需实现
    ``load``，并在其中复用 :meth:`validate_file` 与 :meth:`build_metadata`。
    """

    def __init__(self, file_path: str) -> None:
        """初始化加载器。

        Args:
            file_path: 待加载文件的路径。
        """
        self.file_path = file_path

    @staticmethod
    def validate_file(file_path: str) -> None:
        """校验文件是否存在且为常规文件。

        Args:
            file_path: 文件路径。

        Raises:
            FileNotFoundError: 文件不存在或不是常规文件。
        """
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"文件不存在或不是常规文件: {file_path}")

    @staticmethod
    def build_metadata(file_path: str) -> Dict[str, object]:
        """构建统一的文档元数据。

        Args:
            file_path: 文件路径。

        Returns:
            包含 source / file_name / file_type / page_number / section /
            creation_date / modification_date 的元数据字典。
        """
        return {
            "source": os.path.abspath(file_path),
            "file_name": os.path.basename(file_path),
            "file_type": os.path.splitext(file_path)[1].lower().lstrip("."),
            "page_number": None,
            "section": None,
            "creation_date": datetime.fromtimestamp(os.path.getctime(file_path)).isoformat(),
            "modification_date": datetime.fromtimestamp(os.path.getmtime(file_path)).isoformat(),
        }

    @abstractmethod
    def load(self, file_path: Optional[str] = None) -> List[Document]:
        """加载文档并返回 Document 列表。

        Args:
            file_path: 文件路径；为 None 时使用构造时传入的路径。

        Returns:
            加载得到的文档列表。

        Raises:
            LoaderError: 文件解析失败。
        """
        raise NotImplementedError


class PdfLoader(BaseLoader):
    """PDF 文档加载器。

    基于 :class:`langchain_community.document_loaders.PyPDFLoader` 实现，
    按页拆分文档，并为每一页记录页码（``page_number``）。
    """

    def load(self, file_path: Optional[str] = None) -> List[Document]:
        file_path = file_path or self.file_path
        self.validate_file(file_path)
        try:
            from langchain_community.document_loaders import PyPDFLoader

            raw_documents = PyPDFLoader(file_path).load()
        except Exception as exc:
            logger.error("PDF 文件解析失败: %s（%s）", file_path, exc)
            raise LoaderError(f"PDF 文件解析失败: {file_path}: {exc}") from exc

        documents: List[Document] = []
        for page_number, raw_doc in enumerate(raw_documents, start=1):
            metadata = self.build_metadata(file_path)
            metadata["page_number"] = page_number
            documents.append(
                Document(page_content=raw_doc.page_content, metadata={**raw_doc.metadata, **metadata})
            )
        logger.info("成功加载 PDF 文件: %s，共 %d 页", file_path, len(documents))
        return documents


class DocxLoader(BaseLoader):
    """Word(.docx) 文档加载器。

    基于 python-docx 读取正文段落文本，整体作为单个 Document 返回。
    """

    def load(self, file_path: Optional[str] = None) -> List[Document]:
        file_path = file_path or self.file_path
        self.validate_file(file_path)
        try:
            from docx import Document as WordDocument

            word_document = WordDocument(file_path)
            paragraphs = [p.text.strip() for p in word_document.paragraphs if p.text.strip()]
        except Exception as exc:
            logger.error("Word 文档解析失败: %s（%s）", file_path, exc)
            raise LoaderError(f"Word 文档解析失败: {file_path}: {exc}") from exc

        if not paragraphs:
            logger.warning("Word 文档无可提取的文本段落，返回空结果: %s", file_path)
            return []
        content = "\n".join(paragraphs)
        logger.info("成功加载 Word 文档: %s，共 %d 段", file_path, len(paragraphs))
        return [Document(page_content=content, metadata=self.build_metadata(file_path))]


class MarkdownLoader(BaseLoader):
    """Markdown(.md) 文档加载器。

    直接以 UTF-8 编码读取文件全文，整体作为单个 Document 返回。
    """

    def load(self, file_path: Optional[str] = None) -> List[Document]:
        file_path = file_path or self.file_path
        self.validate_file(file_path)
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
        except (OSError, UnicodeDecodeError) as exc:
            logger.error("Markdown 文件读取失败: %s（%s）", file_path, exc)
            raise LoaderError(f"Markdown 文件读取失败: {file_path}: {exc}") from exc

        if not content.strip():
            logger.warning("Markdown 文件内容为空，返回空结果: %s", file_path)
            return []
        logger.info("成功加载 Markdown 文件: %s，共 %d 字符", file_path, len(content))
        return [Document(page_content=content, metadata=self.build_metadata(file_path))]


class TxtLoader(BaseLoader):
    """TXT 纯文本加载器。

    基于 :class:`langchain_community.document_loaders.TextLoader` 实现，
    以 UTF-8 编码读取文件，整体作为单个 Document 返回。
    """

    def load(self, file_path: Optional[str] = None) -> List[Document]:
        file_path = file_path or self.file_path
        self.validate_file(file_path)
        try:
            raw_documents = _LangchainTextLoader(file_path, encoding="utf-8").load()
        except Exception as exc:
            logger.error("TXT 文件解析失败: %s（%s）", file_path, exc)
            raise LoaderError(f"TXT 文件解析失败: {file_path}: {exc}") from exc

        documents = [
            Document(page_content=doc.page_content, metadata=self.build_metadata(file_path))
            for doc in raw_documents
            if doc.page_content and doc.page_content.strip()
        ]
        if not documents:
            logger.warning("TXT 文件内容为空，返回空结果: %s", file_path)
            return []
        logger.info("成功加载 TXT 文件: %s，共 %d 个文档块", file_path, len(documents))
        return documents


class LoaderFactory:
    """文档加载器工厂。

    通过 :meth:`get_loader` 根据文件扩展名分发到具体的加载器实现。
    """

    _LOADERS: Dict[str, Type[BaseLoader]] = {
        ".pdf": PdfLoader,
        ".docx": DocxLoader,
        ".md": MarkdownLoader,
        ".txt": TxtLoader,
    }

    @staticmethod
    def get_loader(file_path: str) -> BaseLoader:
        """根据文件扩展名返回对应的加载器实例。

        Args:
            file_path: 待加载文件的路径。

        Returns:
            与文件格式匹配的加载器实例。

        Raises:
            FileNotFoundError: 文件不存在。
            UnsupportedFormatError: 文件扩展名不受支持。
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        extension = os.path.splitext(file_path)[1].lower()
        loader_class = LoaderFactory._LOADERS.get(extension)
        if loader_class is None:
            raise UnsupportedFormatError(file_path, extension)

        logger.info("为文件 %s 创建 %s 加载器", file_path, loader_class.__name__)
        return loader_class(file_path)
